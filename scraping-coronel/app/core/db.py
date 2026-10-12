import os
import sqlite3
import openpyxl
from datetime import datetime
from contextlib import closing
from typing import List, Dict, Optional, Any
from app.core.models import Producto

def get_default_db_path() -> str:
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    return os.path.join(base_dir, 'productos.db')

def asegurar_tabla_productos(db_path: Optional[str] = None):
    """
    Garantiza que la tabla 'productos' exista con todas las columnas necesarias
    para el soporte de precios (costo, minorista, mayorista) y jerarquía de categorías.
    Realiza auto-migración segura con ALTER TABLE si faltan columnas.
    """
    if not db_path:
        db_path = get_default_db_path()

    os.makedirs(os.path.dirname(os.path.abspath(db_path)), exist_ok=True)

    with closing(sqlite3.connect(db_path)) as conn:
        cursor = conn.cursor()
        cursor.execute('''
            CREATE TABLE IF NOT EXISTS productos (
                codigo TEXT PRIMARY KEY,
                codigo_barra TEXT,
                descripcion TEXT,
                precio TEXT,
                precio_costo REAL,
                margen_minorista REAL DEFAULT 40.0,
                margen_mayorista REAL DEFAULT 18.0,
                precio_minorista REAL,
                precio_mayorista REAL,
                precio_venta REAL,
                porcentaje_ganancia REAL,
                imagen_url TEXT,
                imagen_local TEXT,
                foto TEXT,
                variante TEXT,
                categoria TEXT,
                subcategoria TEXT,
                sub_subcategoria TEXT,
                categoria_ruta TEXT,
                categoria_slug TEXT,
                subcategoria_slug TEXT,
                sub_subcategoria_slug TEXT,
                stock INTEGER DEFAULT 1,
                peso_kg TEXT,
                ancho_cm TEXT,
                alto_cm TEXT,
                profundidad_cm TEXT,
                updated_at TEXT
            )
        ''')

        # Auto-migración si faltan columnas en una base de datos existente
        cursor.execute("PRAGMA table_info(productos)")
        existentes = {col[1] for col in cursor.fetchall()}

        columnas_a_verificar = [
            ("precio_costo", "REAL"),
            ("margen_minorista", "REAL DEFAULT 40.0"),
            ("margen_mayorista", "REAL DEFAULT 18.0"),
            ("precio_minorista", "REAL"),
            ("precio_mayorista", "REAL"),
            ("precio_venta", "REAL"),
            ("porcentaje_ganancia", "REAL"),
            ("foto", "TEXT"),
            ("sub_subcategoria", "TEXT"),
            ("categoria_ruta", "TEXT"),
            ("categoria_slug", "TEXT"),
            ("subcategoria_slug", "TEXT"),
            ("sub_subcategoria_slug", "TEXT"),
            ("stock", "INTEGER DEFAULT 1"),
            ("updated_at", "TEXT")
        ]

        for col_name, col_type in columnas_a_verificar:
            if col_name not in existentes:
                try:
                    cursor.execute(f"ALTER TABLE productos ADD COLUMN {col_name} {col_type}")
                except Exception:
                    pass

        # Backfill automático de registros preexistentes que tenían precio como string
        try:
            cursor.execute("SELECT codigo, precio, categoria, subcategoria, sub_subcategoria FROM productos WHERE precio_costo IS NULL AND precio IS NOT NULL")
            filas = cursor.fetchall()
            for r_cod, r_prec, r_cat, r_sub, r_subsub in filas:
                if not r_prec:
                    continue
                try:
                    p_limpio = str(r_prec).replace('$', '').replace(' ', '').strip()
                    if ',' in p_limpio and '.' in p_limpio:
                        p_limpio = p_limpio.replace('.', '').replace(',', '.')
                    elif ',' in p_limpio:
                        p_limpio = p_limpio.replace(',', '.')
                    costo_val = float(p_limpio)
                    p_min = round(costo_val * 1.40, 2)
                    p_may = round(costo_val * 1.18, 2)
                    ruta = " > ".join([p for p in [r_cat, r_sub, r_subsub] if p])
                    cursor.execute("""
                        UPDATE productos SET
                            precio_costo = ?,
                            precio_minorista = ?,
                            precio_mayorista = ?,
                            precio_venta = ?,
                            porcentaje_ganancia = 40.0,
                            categoria_ruta = COALESCE(categoria_ruta, ?)
                        WHERE codigo = ?
                    """, (costo_val, p_min, p_may, p_min, ruta, r_cod))
                except Exception:
                    pass
        except Exception:
            pass

        conn.commit()

def inicializar_bd(excel_path: str = 'productos_coronel', db_path: Optional[str] = None) -> Optional[str]:
    """Inicializa la base de datos y cruza códigos de barra del Excel más reciente si existe."""
    try:
        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
        excel_dir = os.path.join(base_dir, excel_path)
        if not db_path:
            db_path = get_default_db_path()

        asegurar_tabla_productos(db_path)

        if not os.path.exists(excel_dir):
            return db_path

        # Buscar Excel más reciente
        archivos = [f for f in os.listdir(excel_dir) if f.lower().endswith('.xlsx')]
        if not archivos:
            return db_path

        archivo_reciente = max(archivos, key=lambda f: os.path.getmtime(os.path.join(excel_dir, f)))
        ruta_excel = os.path.join(excel_dir, archivo_reciente)

        with closing(sqlite3.connect(db_path)) as conn:
            cursor = conn.cursor()
            wb = openpyxl.load_workbook(ruta_excel)
            ws = wb.active

            for row in range(2, ws.max_row + 1):
                codigo = str(ws.cell(row=row, column=1).value or "").strip().replace(" ", "")
                codigo_barra = str(ws.cell(row=row, column=3).value or "").strip()

                if codigo:
                    cursor.execute(
                        "INSERT OR IGNORE INTO productos (codigo, codigo_barra) VALUES (?, ?)",
                        (codigo, codigo_barra)
                    )
                    if codigo_barra:
                        cursor.execute(
                            "UPDATE productos SET codigo_barra = ? WHERE codigo = ?",
                            (codigo_barra, codigo)
                        )

            conn.commit()
            return db_path

    except Exception as e:
        print(f"Advertencia inicializando SQLite desde Excel: {str(e)}")
        if db_path:
            asegurar_tabla_productos(db_path)
            return db_path
        return None

def obtener_codigo_barra(code: str, db_path: Optional[str] = None) -> str:
    """Consulta rápida a SQLite para obtener el código de barras."""
    if not db_path:
        db_path = get_default_db_path()
    try:
        with closing(sqlite3.connect(db_path)) as conn:
            cursor = conn.cursor()
            cursor.execute(
                "SELECT codigo_barra FROM productos WHERE codigo = ?",
                (code.strip(),)
            )
            result = cursor.fetchone()
            return result[0] if result and result[0] else ""
    except Exception as e:
        return ""

def consolidar_productos(todos_los_productos: List[Dict], db_path: Optional[str] = None):
    """
    Guarda todos los datos de productos en SQLite con soporte completo
    de precios originales, márgenes B2B/B2C y jerarquía de 3 niveles de categorías.
    """
    if not db_path:
        db_path = get_default_db_path()

    asegurar_tabla_productos(db_path)

    try:
        with closing(sqlite3.connect(db_path)) as conn:
            cursor = conn.cursor()

            for producto in todos_los_productos:
                codigo = producto.get('codigo', '')
                if not codigo:
                    continue

                codigo_barra = producto.get('codigo_de_barras', '')
                cursor.execute("INSERT OR IGNORE INTO productos (codigo, codigo_barra) VALUES (?, ?)", (codigo, codigo_barra))

                # Extraer o calcular precios y márgenes
                precio_costo = float(producto.get('precio_costo') or 0.0)
                if precio_costo == 0.0 and producto.get('precio'):
                    # Fallback parseo de precio string
                    try:
                        p_limpio = str(producto['precio']).replace('$', '').replace(' ', '')
                        if ',' in p_limpio and '.' in p_limpio:
                            p_limpio = p_limpio.replace('.', '').replace(',', '.')
                        elif ',' in p_limpio:
                            p_limpio = p_limpio.replace(',', '.')
                        precio_costo = float(p_limpio)
                    except Exception:
                        precio_costo = 0.0

                margen_minorista = float(producto.get('margen_minorista') or 40.0)
                margen_mayorista = float(producto.get('margen_mayorista') or 18.0)

                precio_minorista = float(producto.get('precio_minorista') or round(precio_costo * (1.0 + margen_minorista / 100.0), 2))
                precio_mayorista = float(producto.get('precio_mayorista') or round(precio_costo * (1.0 + margen_mayorista / 100.0), 2))

                # Categorías y jerarquía
                categoria = producto.get('categoria', '')
                subcategoria = producto.get('subcategoria', '')
                sub_subcategoria = producto.get('sub_subcategoria', '')

                categoria_ruta = producto.get('categoria_ruta', '')
                if not categoria_ruta:
                    partes = [p for p in [categoria, subcategoria, sub_subcategoria] if p]
                    categoria_ruta = " > ".join(partes)

                categoria_slug = producto.get('categoria_slug', '')
                subcategoria_slug = producto.get('subcategoria_slug', '')
                sub_subcategoria_slug = producto.get('sub_subcategoria_slug', '')

                foto = producto.get('imagen_local') or producto.get('imagen_url') or ""
                ahora = datetime.now().isoformat()

                cursor.execute("""
                    UPDATE productos SET
                        codigo_barra = COALESCE(NULLIF(?, ''), codigo_barra),
                        descripcion = ?,
                        precio = ?,
                        precio_costo = ?,
                        margen_minorista = ?,
                        margen_mayorista = ?,
                        precio_minorista = ?,
                        precio_mayorista = ?,
                        precio_venta = ?,
                        porcentaje_ganancia = ?,
                        imagen_url = ?,
                        imagen_local = ?,
                        foto = ?,
                        variante = ?,
                        categoria = ?,
                        subcategoria = ?,
                        sub_subcategoria = ?,
                        categoria_ruta = ?,
                        categoria_slug = ?,
                        subcategoria_slug = ?,
                        sub_subcategoria_slug = ?,
                        stock = 1,
                        updated_at = ?
                    WHERE codigo = ?
                """, (
                    codigo_barra,
                    producto.get('descripcion', ''),
                    producto.get('precio', ''),
                    precio_costo,
                    margen_minorista,
                    margen_mayorista,
                    precio_minorista,
                    precio_mayorista,
                    precio_minorista,  # alias precio_venta
                    margen_minorista,  # alias porcentaje_ganancia
                    producto.get('imagen_url', ''),
                    producto.get('imagen_local', ''),
                    foto,
                    producto.get('variante', ''),
                    categoria,
                    subcategoria,
                    sub_subcategoria,
                    categoria_ruta,
                    categoria_slug,
                    subcategoria_slug,
                    sub_subcategoria_slug,
                    ahora,
                    codigo
                ))

            # Eliminar productos con descripción vacía
            cursor.execute("DELETE FROM productos WHERE descripcion = '' OR descripcion IS NULL")
            conn.commit()
            print(f"✔ Se consolidaron {len(todos_los_productos)} productos en SQLite con precios y jerarquía completos")
            return todos_los_productos

    except Exception as e:
        print(f"Error consolidando productos en SQLite: {str(e)}")
        return False

def actualizar_margenes_masivo(
    margen_minorista: Optional[float] = None,
    margen_mayorista: Optional[float] = None,
    categoria: Optional[str] = None,
    subcategoria: Optional[str] = None,
    db_path: Optional[str] = None
) -> Dict[str, Any]:
    """
    Recalcula masivamente los precios de venta en base a nuevos márgenes porcentuales
    sin necesidad de re-scrapear.
    """
    if not db_path:
        db_path = get_default_db_path()

    asegurar_tabla_productos(db_path)

    where_clauses = ["precio_costo IS NOT NULL", "precio_costo > 0"]
    params = []

    if categoria:
        where_clauses.append("UPPER(categoria) = ?")
        params.append(categoria.strip().upper())
    if subcategoria:
        where_clauses.append("UPPER(subcategoria) = ?")
        params.append(subcategoria.strip().upper())

    where_sql = " AND ".join(where_clauses)

    with closing(sqlite3.connect(db_path)) as conn:
        cursor = conn.cursor()

        updates = []
        update_params = []

        if margen_minorista is not None:
            updates.append("margen_minorista = ?")
            update_params.append(float(margen_minorista))
            updates.append("precio_minorista = ROUND(precio_costo * (1.0 + ? / 100.0), 2)")
            update_params.append(float(margen_minorista))
            updates.append("precio_venta = ROUND(precio_costo * (1.0 + ? / 100.0), 2)")
            update_params.append(float(margen_minorista))
            updates.append("porcentaje_ganancia = ?")
            update_params.append(float(margen_minorista))

        if margen_mayorista is not None:
            updates.append("margen_mayorista = ?")
            update_params.append(float(margen_mayorista))
            updates.append("precio_mayorista = ROUND(precio_costo * (1.0 + ? / 100.0), 2)")
            update_params.append(float(margen_mayorista))

        if not updates:
            return {"updated": 0, "message": "No se especificaron márgenes a actualizar"}

        updates.append("updated_at = ?")
        update_params.append(datetime.now().isoformat())

        full_sql = f"UPDATE productos SET {', '.join(updates)} WHERE {where_sql}"
        cursor.execute(full_sql, update_params + params)
        affected = cursor.rowcount
        conn.commit()
        return {"updated": affected, "success": True}

def obtener_productos_de_db(db_path: Optional[str] = None) -> List[Dict]:
    """Obtiene todos los productos de SQLite como diccionarios con todos los atributos."""
    if not db_path:
        db_path = get_default_db_path()

    asegurar_tabla_productos(db_path)

    try:
        with closing(sqlite3.connect(db_path)) as conn:
            conn.row_factory = sqlite3.Row
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM productos")
            rows = cursor.fetchall()
            return [dict(row) for row in rows]
    except Exception as e:
        print(f"Error leyendo productos de SQLite: {str(e)}")
        return []

def obtener_productos_objetos(db_path: Optional[str] = None) -> List[Producto]:
    """Obtiene todos los productos de SQLite como objetos Pydantic Producto fuertemente tipados."""
    if not db_path:
        db_path = get_default_db_path()

    asegurar_tabla_productos(db_path)

    try:
        with closing(sqlite3.connect(db_path)) as conn:
            cursor = conn.cursor()
            cursor.execute("SELECT * FROM productos")
            rows = cursor.fetchall()
            return [Producto.from_row(row) for row in rows]
    except Exception as e:
        print(f"Error leyendo productos tipados de SQLite: {str(e)}")
        return []

def actualizar_dimensiones_en_bd(productos: List[Dict], dimensiones: List[Dict], db_path: Optional[str] = None):
    """Actualiza la base de datos SQLite con las dimensiones de envío."""
    if not db_path:
        db_path = get_default_db_path()

    try:
        with closing(sqlite3.connect(db_path)) as conn:
            cursor = conn.cursor()
            for producto, dim in zip(productos, dimensiones):
                cursor.execute("""
                    UPDATE productos SET
                        peso_kg = ?,
                        ancho_cm = ?,
                        alto_cm = ?,
                        profundidad_cm = ?
                    WHERE codigo = ?
                """, (
                    dim.get('peso_kg', ''),
                    dim.get('ancho_cm', ''),
                    dim.get('alto_cm', ''),
                    dim.get('profundidad_cm', ''),
                    producto['codigo']
                ))
            conn.commit()
            print(f"Actualizadas dimensiones para {len(dimensiones)} productos")
    except Exception as e:
        print(f"Error actualizando dimensiones en SQLite: {str(e)}")
